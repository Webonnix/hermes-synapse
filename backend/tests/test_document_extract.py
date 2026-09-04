"""Coverage for backend/document_extract.py: text, PDF, DOCX and spreadsheet
extraction, and the prompt built out of whatever was (or wasn't) extracted --
shared by bot.py's and agent_bot.py's Telegram document handlers."""

from __future__ import annotations

import io
from unittest.mock import MagicMock, patch

import pytest

from backend import document_extract


def test_plain_text_file_is_decoded():
    data = "Привет, это заметка.".encode("utf-8")
    assert document_extract.extract_text(data, "note.txt") == "Привет, это заметка."


def test_cp1251_fallback_when_utf8_fails():
    text = "Старый файл в кодировке cp1251."
    data = text.encode("cp1251")
    assert document_extract.extract_text(data, "old.txt") == text


def test_unsupported_binary_format_returns_none():
    assert document_extract.extract_text(b"\x00\x01binary-junk", "archive.zip") is None


def test_pdf_extraction_concatenates_pages():
    fake_reader = MagicMock()
    page1, page2 = MagicMock(), MagicMock()
    page1.extract_text.return_value = "Page one text."
    page2.extract_text.return_value = "Page two text."
    fake_reader.pages = [page1, page2]

    with patch("pypdf.PdfReader", return_value=fake_reader):
        result = document_extract.extract_text(b"%PDF-fake-bytes", "report.pdf")

    assert result == "Page one text.\n\nPage two text."


def test_pdf_extraction_failure_returns_none_not_raises():
    with patch("pypdf.PdfReader", side_effect=Exception("corrupt file")):
        assert document_extract.extract_text(b"not-a-real-pdf", "broken.pdf") is None


def test_docx_extraction_reads_paragraphs_and_tables():
    import docx

    document = docx.Document()
    document.add_paragraph("Первый абзац.")
    document.add_paragraph("Второй абзац.")
    table = document.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Ключ"
    table.rows[0].cells[1].text = "Значение"
    buffer = io.BytesIO()
    document.save(buffer)

    result = document_extract.extract_text(buffer.getvalue(), "report.docx")

    assert "Первый абзац." in result
    assert "Второй абзац." in result
    assert "Ключ | Значение" in result


def test_xlsx_extraction_reads_sheet_as_csv():
    pd = pytest.importorskip("pandas")

    buffer = io.BytesIO()
    pd.DataFrame({"Имя": ["Аня", "Боря"], "Возраст": [30, 25]}).to_excel(buffer, index=False, sheet_name="Люди")
    buffer.seek(0)

    result = document_extract.extract_text(buffer.getvalue(), "people.xlsx")

    assert "Люди" in result
    assert "Аня" in result and "30" in result
    assert "Боря" in result and "25" in result


def test_build_document_prompt_uses_extracted_text_and_caption():
    data = "Заметка о встрече.".encode("utf-8")
    prompt = document_extract.build_document_prompt("meeting.txt", "text/plain", len(data), data, caption="Что тут?")
    assert prompt.startswith("Что тут?")
    assert "Заметка о встрече." in prompt


def test_build_document_prompt_notes_unreadable_format_without_crashing():
    data = b"\x00binary"
    prompt = document_extract.build_document_prompt("archive.zip", "application/zip", len(data), data)
    assert "archive.zip" in prompt
    assert "application/zip" in prompt


def test_build_document_prompt_truncates_long_text():
    data = ("x" * (document_extract.MAX_EXTRACT_CHARS + 500)).encode("utf-8")
    prompt = document_extract.build_document_prompt("big.txt", "text/plain", len(data), data)
    assert "обрезан" in prompt
    assert len(prompt) < len(data) + 500
